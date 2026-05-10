from __future__ import annotations

import csv
import html
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Iterator, Optional


ROOT = Path(__file__).resolve().parents[2]
DOC_ROOT = ROOT / "uav" / "20-dev-docs"
OUT_ROOT = ROOT / "docs"
SCRIPT_ROOT = Path(__file__).resolve().parent


def candidate_runtime_roots() -> Iterator[Path]:
    for env_name in ("UAV_DOCS_RUNTIME_ROOT", "CODEX_RUNTIME_DEPS"):
        raw_path = os.environ.get(env_name)
        if raw_path:
            yield Path(raw_path).expanduser()

    yield Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies"


def python_bundle_roots() -> Iterator[Path]:
    raw_path = os.environ.get("UAV_DOCS_PYTHON_BUNDLE")
    if raw_path:
        yield Path(raw_path).expanduser()

    for runtime_root in candidate_runtime_roots():
        yield runtime_root / "python"


def node_executable() -> Path:
    raw_path = os.environ.get("UAV_DOCS_NODE")
    if raw_path:
        return Path(raw_path).expanduser()

    for runtime_root in candidate_runtime_roots():
        candidate = runtime_root / "node" / "bin" / "node"
        if candidate.exists():
            return candidate

    resolved = shutil.which("node")
    if resolved:
        return Path(resolved)

    raise RuntimeError(
        "node executable not found; install node or set UAV_DOCS_RUNTIME_ROOT/UAV_DOCS_NODE"
    )


def node_modules_path() -> Optional[Path]:
    raw_path = os.environ.get("UAV_DOCS_NODE_MODULES")
    if raw_path:
        return Path(raw_path).expanduser()

    for runtime_root in candidate_runtime_roots():
        candidate = runtime_root / "node" / "node_modules"
        if candidate.exists():
            return candidate

    for candidate in (SCRIPT_ROOT / "node_modules", ROOT / "node_modules"):
        if candidate.exists():
            return candidate

    return None


def ensure_docx_import() -> None:
    try:
        import docx  # noqa: F401
        return
    except ModuleNotFoundError:
        pass

    for bundle_root in python_bundle_roots():
        for path in sorted(bundle_root.glob("lib/python*/site-packages")):
            path_text = str(path)
            if path_text not in sys.path:
                sys.path.append(path_text)

    try:
        import docx  # noqa: F401
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "python-docx not found; install python-docx or set UAV_DOCS_RUNTIME_ROOT/"
            "UAV_DOCS_PYTHON_BUNDLE"
        ) from exc


def ensure_node_modules_link(node_modules: Optional[Path]) -> tuple[Optional[Path], bool]:
    if node_modules is None:
        return None, False

    link_path = SCRIPT_ROOT / "node_modules"
    node_modules = node_modules.resolve()

    if link_path.is_symlink():
        if link_path.resolve() == node_modules:
            return link_path, False
        link_path.unlink()
    elif link_path.exists():
        if link_path.resolve() == node_modules:
            return link_path, False
        raise RuntimeError(f"unexpected node_modules path exists: {link_path}")

    os.symlink(node_modules, link_path, target_is_directory=True)
    return link_path, True


def load_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def build_export_payload() -> dict[str, object]:
    return {
        "interface_rows": load_csv_rows(DOC_ROOT / "02-interface-table.csv"),
        "feature_rows": load_csv_rows(DOC_ROOT / "03-feature-deployment-matrix.csv"),
        "gap_rows": load_csv_rows(DOC_ROOT / "07-material-gaps.csv"),
        "evidence_rows": load_csv_rows(DOC_ROOT / "06-evidence-index.csv"),
        "guide_markdown": (DOC_ROOT / "01-developer-guide.md").read_text(encoding="utf-8"),
        "appendix_markdown": (DOC_ROOT / "04-screenshot-state-appendix.md").read_text(encoding="utf-8"),
        "navigation_markdown": (DOC_ROOT / "05-developer-navigation.md").read_text(encoding="utf-8"),
    }


def render_inline_markup(text: str) -> str:
    parts = re.split(r"(`[^`]+`)", text)
    rendered: list[str] = []
    for part in parts:
        if part.startswith("`") and part.endswith("`") and len(part) >= 2:
            rendered.append(f"<code>{html.escape(part[1:-1])}</code>")
        else:
            rendered.append(html.escape(part))
    return "".join(rendered)


def markdown_to_html(markdown_text: str) -> str:
    output: list[str] = []
    list_mode: str | None = None

    def close_list() -> None:
        nonlocal list_mode
        if list_mode == "ul":
            output.append("</ul>")
        elif list_mode == "ol":
            output.append("</ol>")
        list_mode = None

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            close_list()
            continue

        if stripped.startswith("# "):
            close_list()
            output.append(f"<h1>{render_inline_markup(stripped[2:])}</h1>")
            continue
        if stripped.startswith("## "):
            close_list()
            output.append(f"<h2>{render_inline_markup(stripped[3:])}</h2>")
            continue
        if stripped.startswith("### "):
            close_list()
            output.append(f"<h3>{render_inline_markup(stripped[4:])}</h3>")
            continue
        if re.match(r"^\d+\.\s+", stripped):
            if list_mode != "ol":
                close_list()
                output.append("<ol>")
                list_mode = "ol"
            content = re.sub(r"^\d+\.\s+", "", stripped)
            output.append(f"<li>{render_inline_markup(content)}</li>")
            continue
        if stripped.startswith("- "):
            if list_mode != "ul":
                close_list()
                output.append("<ul>")
                list_mode = "ul"
            output.append(f"<li>{render_inline_markup(stripped[2:])}</li>")
            continue

        close_list()
        css_class = "chain-line" if stripped.startswith("->") or stripped.startswith("`run_") else "body-line"
        output.append(f'<p class="{css_class}">{render_inline_markup(stripped)}</p>')

    close_list()
    return "\n".join(output)


def feature_status_text(row: dict[str, str]) -> str:
    if row["onboard_deployed"] == "yes":
        return "当前已部署"
    if row["repo_source_present"] == "partial":
        return "源码片段存在，但非实机主链"
    if row["repo_source_present"] == "yes":
        return "源码存在，但未作为当前主链"
    return "产品线文档提到，但当前机上未落地"


def build_html_portal(payload: dict[str, object], out_path: Path) -> None:
    interface_rows = payload["interface_rows"]
    feature_rows = payload["feature_rows"]
    gap_rows = payload["gap_rows"]
    evidence_rows = payload["evidence_rows"]
    guide_markdown = payload["guide_markdown"]
    appendix_markdown = payload["appendix_markdown"]
    navigation_markdown = payload["navigation_markdown"]

    interface_by_name = {row["name"]: row for row in interface_rows}
    evidence_by_id = {row["evidence_id"]: row for row in evidence_rows}
    deployed_rows = [row for row in feature_rows if row["onboard_deployed"] == "yes"]
    extension_rows = [row for row in feature_rows if row["onboard_deployed"] != "yes"]

    recommended_names = [
        "/mavros/state",
        "/mavros/battery",
        "/mavros/rc/in",
        "/ekf/ekf_odom",
        "/vins/imu_propagate",
        "/px4ctrl/takeoff_land",
        "/move_base_simple/goal",
        "/back_trigger",
    ]
    recommended_cards = [interface_by_name[name] for name in recommended_names if name in interface_by_name]

    scenario_rows = [
        {
            "title": "看电量",
            "summary": "先确认 mavros 在线，再看电压和 percentage。",
            "keys": ["/mavros/battery", "/mavros/state", "/mavros/rc/in"],
            "next_step": "电量异常时顺手联看飞控状态和遥控器输入，避免把供电问题看成规划问题。",
            "query": "battery",
            "filter": "state",
        },
        {
            "title": "看飞控状态",
            "summary": "先看 connected / armed / mode，别直接猜 PX4 当前状态。",
            "keys": ["/mavros/state"],
            "next_step": "起飞前、接管异常、动作无响应时都先看它。",
            "query": "/mavros/state",
            "filter": "state",
        },
        {
            "title": "看定位是否正常",
            "summary": "按当前启动链二选一，不要把 LIO 和 VIO 混着判断。",
            "keys": ["/ekf/ekf_odom", "/vins/imu_propagate", "/laserMapping/cloud_registered"],
            "next_step": "LIO 看 /ekf/ekf_odom + 点云，VIO 看 /vins/imu_propagate + 相机链。",
            "query": "odom",
            "filter": "state",
        },
        {
            "title": "起飞 / 开始任务 / 返程 / 降落",
            "summary": "优先用现成动作原语，不要直接从控制器最底层切入。",
            "keys": ["/px4ctrl/takeoff_land", "/move_base_simple/goal", "/back_trigger"],
            "next_step": "触发动作前，确认 RC Channel 8 不会和脚本触发打架。",
            "query": "",
            "filter": "action",
        },
        {
            "title": "改航点 / 起飞高度",
            "summary": "先改参数文件，不要误改到底层规划或控制器核心。",
            "keys": ["points.yaml", "ctrl_param_fpv.yaml", "multipointplan_exp_lio.launch"],
            "next_step": "航点看 points.yaml，起飞高度看 ctrl_param_fpv.yaml，任务触发行为看 multipointplan。",
            "query": "",
            "filter": "parameter",
        },
        {
            "title": "接上层 Agent / MQTT",
            "summary": "第一刀先做状态订阅 + 动作发布闭环，先把接口边界站稳。",
            "keys": ["/mavros/state", "/mavros/battery", "/px4ctrl/takeoff_land", "/move_base_simple/goal"],
            "next_step": "先接状态观测，再补动作原语；第一轮不要直接碰 /setpoints_cmd 和 px4ctrl 状态机。",
            "query": "mavros",
            "filter": "state",
        },
    ]

    startup_chain_rows = [
        {
            "title": "LIO 全栈启动",
            "summary": "机上雷达定位主链。雷达、融合、规划、控制和任务编排都在同一个入口脚本里拉起。",
            "entrypoint": "run_single_lio.sh",
            "steps": ["run_single_lio.sh", "mavros", "faster_lio", "ekf", "planner", "px4ctrl", "multipoint"],
            "checks": ["/mavros/state", "/ekf/ekf_odom", "/laserMapping/cloud_registered"],
            "risk": "启动较慢时先别误判失败，先确认 ROS graph 还在继续起节点。",
        },
        {
            "title": "VIO 全栈启动",
            "summary": "机上视觉定位主链。D435 与 VINS 是当前视觉路径的核心，适合你后续做视觉相关二开。",
            "entrypoint": "run_single_vio.sh",
            "steps": ["run_single_vio.sh", "mavros", "realsense", "vins", "planner", "px4ctrl", "multipoint"],
            "checks": ["/mavros/state", "/vins/imu_propagate", "相机图像 topic"],
            "risk": "相机或 USB 链路掉了，VIO 整链就失效，不要只盯 planner。",
        },
        {
            "title": "任务动作链",
            "summary": "脚本和 RC 最后都落在动作 topic 上。后续接上层系统时，优先复用这条动作链而不是绕过它。",
            "entrypoint": "takeoff.sh / pub_trigger.sh / back.sh / land.sh",
            "steps": ["动作脚本或 RC 8 通", "动作 topic", "multipoint / px4ctrl", "planner", "/setpoints_cmd", "px4ctrl"],
            "checks": ["/px4ctrl/takeoff_land", "/move_base_simple/goal", "/back_trigger", "/setpoints_cmd"],
            "risk": "动作 topic 和 RC 触发同时存在，二开时必须把仲裁关系想清楚。",
        },
    ]

    machine_cards = [
        ("当前主链", "当前主链是单机 LIO / VIO + Diff-planner + px4ctrl + mavros + multipoint"),
        ("视觉链", "D435 + VINS 是当前视觉定位主链假设"),
        ("当前边界", "Elastic 不属于当前实机默认部署链路"),
        ("主工作空间", "Diff-planner"),
    ]

    def category_label(category: str) -> str:
        return {
            "state": "状态接口",
            "action": "动作接口",
            "script": "启动脚本",
            "parameter": "参数入口",
            "rc": "遥控器映射",
        }.get(category, category)

    def machine_status_text(machine_status: str) -> str:
        if machine_status == "onboard-active":
            return "当前机上主链"
        return machine_status or "未标注"

    def trace_claim_text(row: dict[str, str], evidence: dict[str, str]) -> str:
        claim = evidence.get("claim", "").strip()
        if not claim:
            return f"{row['name']} 用于{row['typical_use']}"

        source_match = row.get("source_path", "") and row.get("source_path", "") == evidence.get("source_path", "")
        keywords = [
            row["name"],
            row["typical_use"],
            row["typical_use"].removeprefix("查看"),
            row["typical_use"].removeprefix("配置"),
            row["typical_use"].removeprefix("定义"),
            row["typical_use"].removeprefix("切换"),
            row["typical_use"].removeprefix("触发"),
        ]
        if source_match or any(keyword and keyword in claim for keyword in keywords):
            return claim
        return f"{row['name']} 用于{row['typical_use']}"

    def modification_layer(row: dict[str, str]) -> str:
        category = row["category"]
        name = row["name"]
        if category == "state":
            return "状态观测层：优先只读订阅，适合监控、日志和上层 Agent 观测。"
        if category == "action":
            if name == "/setpoints_cmd":
                return "规划-控制边界层：这里只适合高级控制改造，不该作为第一接入点。"
            return "动作原语层：优先在这层做任务系统、桥接服务和外部控制接入。"
        if category == "script":
            return "脚本编排层：适合梳理启动链和最小自动化，但不适合长期协议设计。"
        if category == "parameter":
            if name == "ctrl_param_fpv.yaml":
                return "控制配置层：涉及起飞高度、电压阈值和安全边界，改动必须谨慎。"
            if name.startswith("run_exp_single_"):
                return "规划配置层：适合做轨迹、地图和性能实验，不适合日常任务改点。"
            return "任务配置层：第一轮二开优先从这里下手，成本最低，风险也最低。"
        if category == "rc":
            return "人工安全边界层：这是安全兜底，不是普通软件 API。"
        return "待分类层：需要结合源码再细分。"

    def build_trace_payload(row: dict[str, str], *, purpose_override: str | None = None) -> dict[str, str]:
        evidence = evidence_by_id.get(row.get("evidence_id", ""), {})
        return {
            "title": row["name"],
            "kind": category_label(row["category"]),
            "layer": modification_layer(row),
            "purpose": purpose_override or row["typical_use"],
            "evidence": row.get("evidence_id", ""),
            "claim": trace_claim_text(row, evidence),
            "source": row.get("source_path", "") or evidence.get("source_path", ""),
            "preconditions": row.get("preconditions", ""),
            "example": row.get("example", ""),
            "risk": row.get("risks", ""),
            "upstream": row.get("upstream", ""),
            "downstream": row.get("downstream", ""),
            "status": machine_status_text(row.get("machine_status", "")),
            "notes": evidence.get("notes", ""),
        }

    def render_trace_attrs(payload: dict[str, str], *, is_default: bool = False) -> str:
        attrs = {
            "data-trace-title": payload["title"],
            "data-trace-kind": payload["kind"],
            "data-trace-layer": payload["layer"],
            "data-trace-purpose": payload["purpose"],
            "data-trace-evidence": payload["evidence"],
            "data-trace-claim": payload["claim"],
            "data-trace-source": payload["source"],
            "data-trace-preconditions": payload["preconditions"],
            "data-trace-example": payload["example"],
            "data-trace-risk": payload["risk"],
            "data-trace-upstream": payload["upstream"],
            "data-trace-downstream": payload["downstream"],
            "data-trace-status": payload["status"],
            "data-trace-notes": payload["notes"],
        }
        if is_default:
            attrs["data-trace-default"] = "true"
        return " ".join(f'{key}="{html.escape(value)}"' for key, value in attrs.items())

    parameter_focus_specs = [
        (
            "points.yaml",
            "第一手改",
            "任务点和返程点都先从这里改。你要改任务内容，先改它，不要上来就动 planner。",
            "支持 [x,y,z] / [x,y,z,time] / [x,y,z,yaw] / [x,y,z,yaw,time] 四种格式。",
        ),
        (
            "ctrl_param_fpv.yaml",
            "安全参数",
            "起飞高度、低电压阈值、人工接管速度都在这里，属于低频但高风险参数入口。",
            "关注 auto_takeoff_land.takeoff_height、low_voltage、max_manual_vel。",
        ),
        (
            "multipointplan_exp_lio.launch",
            "任务行为",
            "start_plan、back_plan、fligt_type 这类任务行为开关在这里，决定 multipoint 怎么触发和怎么走点。",
            "你如果要改任务模式，不先看它，后面一定会把动作接口理解错。",
        ),
        (
            "run_exp_single_lio.launch",
            "LIO 规划",
            "地图尺寸、速度加速度上限、LIO 规划入口参数都在这里，适合做规划层实验。",
            "这是规划层参数，不是日常任务参数；别拿它当航点配置文件。",
        ),
        (
            "run_exp_single_vio.launch",
            "VIO 规划",
            "VIO 主链对应的规划参数入口。你后续如果主要围绕单目/深度视觉链二开，这个文件必须熟。",
            "改它之前先确认相机链和 VINS 输出已经稳定，否则会把定位问题误判成规划问题。",
        ),
    ]
    parameter_focus_rows: list[dict[str, str]] = []
    for name, badge, summary, focus in parameter_focus_specs:
        if name not in interface_by_name:
            continue
        row = interface_by_name[name]
        parameter_focus_rows.append(
            {
                "category": row["category"],
                "name": name,
                "badge": badge,
                "summary": summary,
                "focus": focus,
                "typical_use": row["typical_use"],
                "key_fields": row["key_fields"],
                "example": row["example"],
                "risks": row["risks"],
                "preconditions": row["preconditions"],
                "upstream": row["upstream"],
                "downstream": row["downstream"],
                "evidence_id": row["evidence_id"],
                "source_path": row["source_path"],
                "machine_status": row["machine_status"],
            }
        )

    nav_items = [
        ("overview", "机器总览"),
        ("scenes", "场景入口"),
        ("deployed", "当前已部署"),
        ("chains", "启动链速览"),
        ("params", "参数聚焦"),
        ("trace", "源码追踪"),
        ("recommended", "推荐接口"),
        ("explorer", "接口检索"),
        ("docs", "完整文档"),
        ("gaps", "素材缺口"),
        ("evidence", "证据索引"),
        ("downloads", "导出件"),
    ]

    def render_pills(items: list[str], css_class: str = "pill") -> str:
        return "".join(f'<span class="{css_class}">{html.escape(item)}</span>' for item in items)

    default_trace_row = interface_by_name.get("points.yaml") or interface_by_name.get("/px4ctrl/takeoff_land") or interface_rows[0]
    default_trace_payload = build_trace_payload(default_trace_row)

    recommended_html = "".join(
        f"""
        <article class="mini-card interface-card trace-trigger" role="button" tabindex="0" data-category="{html.escape(row['category'])}" data-search="{html.escape(' '.join(row.values()))}" {render_trace_attrs(build_trace_payload(row), is_default=row['name'] == default_trace_row['name'])}>
          <div class="card-kicker">{html.escape(category_label(row['category']))}</div>
          <h3>{html.escape(row['name'])}</h3>
          <p>{html.escape(row['typical_use'])}</p>
          <dl class="meta-grid">
            <div><dt>类型</dt><dd>{html.escape(row['type'])}</dd></div>
            <div><dt>方向</dt><dd>{html.escape(row['direction'])}</dd></div>
            <div><dt>前置</dt><dd>{html.escape(row['preconditions'])}</dd></div>
            <div><dt>风险</dt><dd>{html.escape(row['risks'])}</dd></div>
          </dl>
        </article>
        """
        for row in recommended_cards
    )

    explorer_cards = "".join(
        f"""
        <article class="interface-row trace-trigger" role="button" tabindex="0" data-category="{html.escape(row['category'])}" data-search="{html.escape(' '.join(row.values()))}" {render_trace_attrs(build_trace_payload(row), is_default=row['name'] == default_trace_row['name'])}>
          <header>
            <span class="tag">{html.escape(category_label(row['category']))}</span>
            <h3>{html.escape(row['name'])}</h3>
          </header>
          <p class="lead">{html.escape(row['typical_use'])}</p>
          <div class="detail-grid">
            <div><span>类型</span><strong>{html.escape(row['type'])}</strong></div>
            <div><span>方向</span><strong>{html.escape(row['direction'])}</strong></div>
            <div><span>格式</span><strong>{html.escape(row['message_or_format'])}</strong></div>
            <div><span>前置条件</span><strong>{html.escape(row['preconditions'])}</strong></div>
            <div><span>上游</span><strong>{html.escape(row['upstream'])}</strong></div>
            <div><span>下游</span><strong>{html.escape(row['downstream'])}</strong></div>
          </div>
          <p class="card-note"><strong>示例：</strong>{html.escape(row['example'])}</p>
          <p class="card-note danger"><strong>风险：</strong>{html.escape(row['risks'])}</p>
        </article>
        """
        for row in interface_rows
    )

    deployed_html = "".join(
        f"""
        <article class="feature-card deployed">
          <div class="card-kicker">当前已部署</div>
          <h3>{html.escape(row['feature_name'])}</h3>
          <p>{html.escape(row['notes'])}</p>
          <dl class="meta-grid">
            <div><dt>入口</dt><dd>{html.escape(row['primary_entrypoint'])}</dd></div>
            <div><dt>验证</dt><dd>{html.escape(row['verification_method'])}</dd></div>
            <div><dt>硬件</dt><dd>{html.escape(row['required_hardware'])}</dd></div>
            <div><dt>可直接运行</dt><dd>{'是' if row['directly_runnable'] == 'yes' else '否'}</dd></div>
          </dl>
        </article>
        """
        for row in deployed_rows
    )

    extension_html = "".join(
        f"""
        <article class="feature-card extension">
          <div class="card-kicker">扩展能力</div>
          <h3>{html.escape(row['feature_name'])}</h3>
          <p>{html.escape(feature_status_text(row))}</p>
          <p class="card-note">{html.escape(row['notes'])}</p>
          <dl class="meta-grid">
            <div><dt>需要补什么</dt><dd>{html.escape(row['required_hardware'])}</dd></div>
            <div><dt>入口</dt><dd>{html.escape(row['primary_entrypoint'])}</dd></div>
          </dl>
        </article>
        """
        for row in extension_rows
    )

    scenario_html = "".join(
        f"""
        <article class="scenario-card scene-jump" role="button" tabindex="0" data-scene-query="{html.escape(row['query'])}" data-scene-filter="{html.escape(row['filter'])}" data-scene-title="{html.escape(row['title'])}" data-scene-summary="{html.escape(row['next_step'])}">
          <div class="card-kicker">场景入口</div>
          <h3>{html.escape(row['title'])}</h3>
          <p>{html.escape(row['summary'])}</p>
          <div class="pill-row">{render_pills(row['keys'])}</div>
          <p class="card-note">{html.escape(row['next_step'])}</p>
        </article>
        """
        for row in scenario_rows
    )

    startup_html = "".join(
        f"""
        <article class="chain-card">
          <div class="card-kicker">{html.escape(row['entrypoint'])}</div>
          <h3>{html.escape(row['title'])}</h3>
          <p>{html.escape(row['summary'])}</p>
          <div class="flow-row">{render_pills(row['steps'], 'flow-step')}</div>
          <div class="chain-footer">
            <div>
              <strong>先看</strong>
              <div class="pill-row">{render_pills(row['checks'])}</div>
            </div>
            <p class="card-note danger"><strong>风险：</strong>{html.escape(row['risk'])}</p>
          </div>
        </article>
        """
        for row in startup_chain_rows
    )

    parameter_html = "".join(
        f"""
        <article class="param-card trace-trigger" role="button" tabindex="0" {render_trace_attrs(build_trace_payload(row, purpose_override=row['summary']), is_default=row['name'] == default_trace_row['name'])}>
          <div class="param-head">
            <span class="badge">{html.escape(row['badge'])}</span>
            <code>{html.escape(row['name'])}</code>
          </div>
          <h3>{html.escape(row['typical_use'])}</h3>
          <p>{html.escape(row['summary'])}</p>
          <div class="meta-grid">
            <div><dt>先盯这里</dt><dd>{html.escape(row['focus'])}</dd></div>
            <div><dt>前置条件</dt><dd>{html.escape(row['preconditions'])}</dd></div>
            <div><dt>关键字段</dt><dd>{html.escape(row['key_fields'])}</dd></div>
            <div><dt>命令示例</dt><dd>{html.escape(row['example'])}</dd></div>
          </div>
          <p class="card-note danger"><strong>风险：</strong>{html.escape(row['risks'])}</p>
        </article>
        """
        for row in parameter_focus_rows
    )

    gap_html = "".join(
        f"""
        <tr>
          <td>{html.escape(row['state_group'])}</td>
          <td>{html.escape(row['required_shot'])}</td>
          <td>{html.escape(row['gap_reason'])}</td>
          <td>{html.escape(row['collection_instruction'])}</td>
          <td>{html.escape(row['priority'])}</td>
        </tr>
        """
        for row in gap_rows
    )

    evidence_html = "".join(
        f"""
        <tr>
          <td>{html.escape(row['evidence_id'])}</td>
          <td>{html.escape(row['claim'])}</td>
          <td>{html.escape(row['source_path'])}</td>
        </tr>
        """
        for row in evidence_rows
    )

    downloads_html = """
    <div class="download-row">
      <a class="download-link" href="./非凸α-开发者文档包.docx">开发文档包.docx</a>
      <a class="download-link" href="./非凸α-接口总表.xlsx">接口总表.xlsx</a>
      <a class="download-link" href="./非凸α-功能部署矩阵.xlsx">功能部署矩阵.xlsx</a>
      <a class="download-link" href="./非凸α-素材缺口表.xlsx">素材缺口表.xlsx</a>
    </div>
    """

    html_text = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>非凸α二开速查台</title>
  <style>
    :root {{
      --bg: #f3efe5;
      --panel: rgba(255, 252, 245, 0.88);
      --panel-strong: #fffdf7;
      --ink: #13201a;
      --muted: #576258;
      --line: rgba(19, 32, 26, 0.12);
      --accent: #1f6b55;
      --accent-2: #c96a2b;
      --accent-3: #1a3e5c;
      --danger: #8c3b2a;
      --shadow: 0 18px 44px rgba(23, 39, 29, 0.12);
    }}
    * {{ box-sizing: border-box; }}
    html {{ scroll-behavior: smooth; }}
    body {{
      margin: 0;
      font-family: "IBM Plex Sans", "Avenir Next", "PingFang SC", "Hiragino Sans GB", sans-serif;
      color: var(--ink);
      background:
        radial-gradient(circle at top left, rgba(31, 107, 85, 0.16), transparent 32%),
        radial-gradient(circle at 80% 10%, rgba(201, 106, 43, 0.18), transparent 24%),
        linear-gradient(180deg, #f8f3e9 0%, var(--bg) 100%);
      min-height: 100vh;
    }}
    a {{ color: inherit; }}
    code {{
      font-family: "IBM Plex Mono", "SFMono-Regular", Consolas, monospace;
      background: rgba(20, 46, 36, 0.08);
      padding: 0.14rem 0.36rem;
      border-radius: 0.42rem;
      font-size: 0.92em;
    }}
    .layout {{
      display: grid;
      grid-template-columns: 280px minmax(0, 1fr);
      gap: 24px;
      max-width: 1600px;
      margin: 0 auto;
      padding: 24px;
    }}
    .sidebar {{
      position: sticky;
      top: 18px;
      align-self: start;
      padding: 24px 20px;
      border: 1px solid var(--line);
      border-radius: 28px;
      background: rgba(255,255,255,0.52);
      backdrop-filter: blur(14px);
      box-shadow: var(--shadow);
    }}
    .brand {{
      font-family: "Iowan Old Style", "Palatino Linotype", "Songti SC", serif;
      font-size: 1.45rem;
      line-height: 1.15;
      margin: 0 0 10px;
    }}
    .sidebar p {{
      margin: 0 0 18px;
      color: var(--muted);
      font-size: 0.95rem;
      line-height: 1.6;
    }}
    .nav-list {{
      list-style: none;
      padding: 0;
      margin: 0 0 18px;
      display: grid;
      gap: 8px;
    }}
    .nav-list a {{
      display: block;
      padding: 10px 12px;
      border-radius: 14px;
      text-decoration: none;
      color: var(--ink);
      border: 1px solid transparent;
      transition: transform 180ms ease, background 180ms ease, border-color 180ms ease;
    }}
    .nav-list a:hover {{
      transform: translateX(4px);
      background: rgba(31, 107, 85, 0.08);
      border-color: rgba(31, 107, 85, 0.16);
    }}
    .main {{
      display: grid;
      gap: 24px;
    }}
    .hero {{
      padding: 34px;
      border-radius: 34px;
      background:
        linear-gradient(135deg, rgba(255,255,255,0.88), rgba(255,249,240,0.88)),
        linear-gradient(135deg, rgba(31, 107, 85, 0.08), rgba(201, 106, 43, 0.06));
      border: 1px solid var(--line);
      box-shadow: var(--shadow);
      overflow: hidden;
      position: relative;
    }}
    .hero::after {{
      content: "";
      position: absolute;
      inset: auto -80px -100px auto;
      width: 260px;
      height: 260px;
      border-radius: 50%;
      background: radial-gradient(circle, rgba(31, 107, 85, 0.18), transparent 70%);
    }}
    .eyebrow {{
      display: inline-flex;
      align-items: center;
      gap: 8px;
      padding: 8px 12px;
      border-radius: 999px;
      background: rgba(31, 107, 85, 0.1);
      color: var(--accent);
      font-weight: 700;
      letter-spacing: 0.04em;
      text-transform: uppercase;
      font-size: 0.74rem;
    }}
    .hero h1 {{
      font-family: "Iowan Old Style", "Palatino Linotype", "Songti SC", serif;
      font-size: clamp(2.3rem, 5vw, 4rem);
      margin: 14px 0 12px;
      line-height: 0.98;
      letter-spacing: -0.03em;
      max-width: 10ch;
    }}
    .hero .summary {{
      max-width: 78ch;
      font-size: 1.02rem;
      line-height: 1.72;
      color: var(--muted);
      margin: 0 0 24px;
    }}
    .card-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
      gap: 16px;
    }}
    .metric-card, .scenario-card, .feature-card, .mini-card, .doc-card {{
      border-radius: 24px;
      background: var(--panel);
      border: 1px solid var(--line);
      box-shadow: var(--shadow);
      padding: 20px;
    }}
    .trace-trigger {{
      cursor: pointer;
      transition: transform 180ms ease, border-color 180ms ease, background 180ms ease, box-shadow 180ms ease;
    }}
    .trace-trigger:hover,
    .trace-trigger:focus-visible,
    .trace-trigger.active-trace {{
      outline: none;
      transform: translateY(-2px);
      border-color: rgba(31, 107, 85, 0.24);
      background: rgba(255, 255, 255, 0.96);
      box-shadow: 0 22px 42px rgba(23, 39, 29, 0.16);
    }}
    .scenario-card {{
      position: relative;
      cursor: pointer;
      transition: transform 180ms ease, border-color 180ms ease, background 180ms ease, box-shadow 180ms ease;
    }}
    .scenario-card:hover,
    .scenario-card:focus-visible,
    .scenario-card.active-scene {{
      outline: none;
      transform: translateY(-2px);
      border-color: rgba(31, 107, 85, 0.24);
      background: rgba(255, 255, 255, 0.96);
      box-shadow: 0 22px 42px rgba(23, 39, 29, 0.16);
    }}
    .scene-helper,
    .scene-status {{
      padding: 14px 16px;
      border-radius: 18px;
      background: rgba(31, 107, 85, 0.08);
      border: 1px solid rgba(31, 107, 85, 0.14);
      color: var(--ink);
    }}
    .scene-helper {{
      margin-bottom: 18px;
      line-height: 1.65;
      color: var(--muted);
    }}
    .scene-status {{
      display: flex;
      flex-wrap: wrap;
      justify-content: space-between;
      gap: 12px 18px;
      align-items: center;
      margin-bottom: 16px;
    }}
    .scene-status strong {{
      display: block;
      margin-bottom: 2px;
    }}
    .scene-status p {{
      margin: 0;
      color: var(--muted);
      line-height: 1.6;
    }}
    .scene-status-count {{
      color: var(--accent-3);
      font-weight: 700;
      white-space: nowrap;
    }}
    .metric-card strong {{
      display: block;
      font-size: 1.1rem;
      margin-bottom: 8px;
    }}
    .metric-card p,
    .scenario-card p,
    .feature-card p,
    .mini-card p,
    .doc-card p {{
      margin: 0;
      color: var(--muted);
      line-height: 1.68;
    }}
    .section {{
      padding: 26px;
      border-radius: 30px;
      background: rgba(255,255,255,0.56);
      border: 1px solid var(--line);
      box-shadow: var(--shadow);
    }}
    .section-header {{
      display: flex;
      flex-wrap: wrap;
      justify-content: space-between;
      gap: 12px 18px;
      align-items: end;
      margin-bottom: 20px;
    }}
    .section-header h2 {{
      margin: 0;
      font-size: 1.7rem;
      line-height: 1.1;
    }}
    .section-header p {{
      margin: 0;
      max-width: 72ch;
      color: var(--muted);
      line-height: 1.68;
    }}
    .pill-row {{
      display: flex;
      flex-wrap: wrap;
      gap: 8px;
      margin-top: 14px;
    }}
    .pill {{
      display: inline-flex;
      align-items: center;
      padding: 7px 11px;
      border-radius: 999px;
      background: rgba(26, 62, 92, 0.08);
      color: var(--accent-3);
      font-size: 0.88rem;
      border: 1px solid rgba(26, 62, 92, 0.11);
    }}
    .card-kicker {{
      font-size: 0.74rem;
      text-transform: uppercase;
      letter-spacing: 0.08em;
      color: var(--accent);
      font-weight: 700;
      margin-bottom: 8px;
    }}
    .feature-card.deployed .card-kicker {{ color: var(--accent); }}
    .feature-card.extension .card-kicker {{ color: var(--accent-2); }}
    .feature-card h3,
    .scenario-card h3,
    .mini-card h3,
    .interface-row h3 {{
      margin: 0 0 10px;
      font-size: 1.12rem;
      line-height: 1.35;
    }}
    .meta-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(170px, 1fr));
      gap: 12px;
      margin-top: 16px;
    }}
    .meta-grid div {{
      padding: 10px 12px;
      border-radius: 16px;
      background: rgba(19, 32, 26, 0.04);
      border: 1px solid rgba(19, 32, 26, 0.06);
    }}
    .meta-grid dt {{
      margin: 0 0 4px;
      color: var(--muted);
      font-size: 0.8rem;
      text-transform: uppercase;
      letter-spacing: 0.06em;
    }}
    .meta-grid dd {{
      margin: 0;
      line-height: 1.5;
    }}
    .stack-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(280px, 1fr));
      gap: 16px;
    }}
    .chain-card,
    .param-card {{
      border-radius: 24px;
      background: var(--panel);
      border: 1px solid var(--line);
      box-shadow: var(--shadow);
      padding: 20px;
    }}
    .flow-row {{
      display: flex;
      flex-wrap: wrap;
      gap: 10px;
      margin-top: 16px;
    }}
    .flow-step {{
      display: inline-flex;
      align-items: center;
      padding: 9px 12px;
      border-radius: 14px;
      background: rgba(31, 107, 85, 0.08);
      color: var(--accent);
      border: 1px solid rgba(31, 107, 85, 0.12);
      font-family: "IBM Plex Mono", "SFMono-Regular", Consolas, monospace;
      font-size: 0.86rem;
    }}
    .flow-step::after {{
      content: "→";
      margin-left: 10px;
      color: rgba(31, 107, 85, 0.54);
    }}
    .flow-step:last-child::after {{
      display: none;
    }}
    .chain-footer {{
      display: grid;
      gap: 14px;
      margin-top: 18px;
    }}
    .param-head {{
      display: flex;
      flex-wrap: wrap;
      gap: 10px;
      align-items: center;
      margin-bottom: 12px;
    }}
    .badge {{
      display: inline-flex;
      align-items: center;
      padding: 6px 10px;
      border-radius: 999px;
      background: rgba(201, 106, 43, 0.12);
      color: var(--accent-2);
      font-size: 0.82rem;
      font-weight: 700;
    }}
    .trace-layout {{
      display: grid;
      grid-template-columns: minmax(0, 1.25fr) minmax(280px, 0.75fr);
      gap: 16px;
      align-items: start;
    }}
    .trace-panel,
    .trace-guide-card {{
      border-radius: 24px;
      background: var(--panel);
      border: 1px solid var(--line);
      box-shadow: var(--shadow);
      padding: 20px;
    }}
    .trace-panel {{
      position: sticky;
      top: 18px;
    }}
    .trace-guide {{
      display: grid;
      gap: 14px;
    }}
    .trace-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(170px, 1fr));
      gap: 12px;
      margin-top: 16px;
    }}
    .trace-grid div,
    .trace-source,
    .trace-io {{
      padding: 12px 14px;
      border-radius: 18px;
      background: rgba(19, 32, 26, 0.04);
      border: 1px solid rgba(19, 32, 26, 0.06);
    }}
    .trace-grid span,
    .trace-source span,
    .trace-io span {{
      display: block;
      color: var(--muted);
      font-size: 0.82rem;
      margin-bottom: 6px;
    }}
    .trace-source code {{
      display: block;
      white-space: pre-wrap;
      word-break: break-all;
      line-height: 1.6;
    }}
    .trace-io-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
      gap: 12px;
      margin-top: 12px;
    }}
    .controls {{
      display: flex;
      flex-wrap: wrap;
      gap: 10px;
      margin-bottom: 16px;
    }}
    .controls input {{
      flex: 1 1 280px;
      min-width: 220px;
      border: 1px solid var(--line);
      background: var(--panel-strong);
      border-radius: 16px;
      padding: 12px 14px;
      font: inherit;
      color: var(--ink);
    }}
    .filter-buttons {{
      display: flex;
      flex-wrap: wrap;
      gap: 8px;
    }}
    .filter-buttons button,
    .download-link {{
      border: 1px solid var(--line);
      background: var(--panel-strong);
      color: var(--ink);
      border-radius: 999px;
      padding: 10px 14px;
      font: inherit;
      cursor: pointer;
      text-decoration: none;
      transition: transform 160ms ease, background 160ms ease, border-color 160ms ease;
    }}
    .filter-buttons button.active,
    .filter-buttons button:hover,
    .download-link:hover {{
      background: rgba(31, 107, 85, 0.12);
      border-color: rgba(31, 107, 85, 0.18);
      transform: translateY(-1px);
    }}
    .interface-list {{
      display: grid;
      gap: 14px;
    }}
    .interface-row {{
      padding: 18px 20px;
      border-radius: 24px;
      background: var(--panel);
      border: 1px solid var(--line);
    }}
    .interface-row header {{
      display: flex;
      gap: 12px;
      align-items: center;
      margin-bottom: 10px;
      flex-wrap: wrap;
    }}
    .tag {{
      display: inline-flex;
      align-items: center;
      padding: 6px 10px;
      border-radius: 999px;
      background: rgba(201, 106, 43, 0.12);
      color: var(--accent-2);
      font-size: 0.82rem;
      font-weight: 700;
    }}
    .lead {{
      margin: 0 0 14px;
      color: var(--muted);
    }}
    .detail-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(160px, 1fr));
      gap: 10px;
      margin-bottom: 12px;
    }}
    .detail-grid div {{
      padding: 10px 12px;
      border-radius: 16px;
      background: rgba(19, 32, 26, 0.04);
    }}
    .detail-grid span {{
      display: block;
      color: var(--muted);
      font-size: 0.82rem;
      margin-bottom: 4px;
    }}
    .detail-grid strong {{
      font-size: 0.96rem;
      line-height: 1.5;
    }}
    .card-note {{
      margin: 10px 0 0;
      color: var(--muted);
      line-height: 1.65;
    }}
    .card-note.danger {{ color: var(--danger); }}
    .doc-columns {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(320px, 1fr));
      gap: 16px;
    }}
    .doc-card {{
      overflow: hidden;
    }}
    .doc-card h1, .doc-card h2, .doc-card h3 {{
      margin-top: 1.1em;
      margin-bottom: 0.5em;
      line-height: 1.25;
    }}
    .doc-card h1:first-child,
    .doc-card h2:first-child,
    .doc-card h3:first-child {{ margin-top: 0; }}
    .doc-card ul,
    .doc-card ol {{
      padding-left: 1.2rem;
      margin: 0.7rem 0;
      color: var(--muted);
      line-height: 1.68;
    }}
    .doc-card .body-line,
    .doc-card .chain-line {{
      color: var(--muted);
      line-height: 1.68;
      margin: 0.6rem 0;
    }}
    .doc-card .chain-line {{
      font-family: "IBM Plex Mono", "SFMono-Regular", Consolas, monospace;
      color: var(--accent-3);
    }}
    .table-wrap {{
      overflow-x: auto;
      border-radius: 22px;
      border: 1px solid var(--line);
      background: var(--panel);
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      min-width: 760px;
    }}
    th, td {{
      padding: 14px 16px;
      text-align: left;
      vertical-align: top;
      border-bottom: 1px solid var(--line);
      line-height: 1.6;
    }}
    th {{
      background: rgba(31, 107, 85, 0.1);
      font-size: 0.86rem;
      text-transform: uppercase;
      letter-spacing: 0.05em;
    }}
    .download-row {{
      display: flex;
      flex-wrap: wrap;
      gap: 12px;
      margin-top: 18px;
    }}
    .footer-note {{
      color: var(--muted);
      line-height: 1.7;
      margin-top: 10px;
    }}
    @media (max-width: 1080px) {{
      .layout {{
        grid-template-columns: 1fr;
      }}
      .sidebar {{
        position: static;
      }}
      .trace-layout {{
        grid-template-columns: 1fr;
      }}
      .trace-panel {{
        position: static;
      }}
    }}
  </style>
</head>
<body>
  <div class="layout">
    <aside class="sidebar">
      <h1 class="brand">非凸α<br>二开速查台</h1>
      <p>这不是另一本说明书，而是把当前这台机器真正可用的主链、接口、参数、边界和速查入口压进一个本地页面。</p>
      <nav>
        <ul class="nav-list">
          {''.join(f'<li><a href="#{html.escape(anchor)}">{html.escape(label)}</a></li>' for anchor, label in nav_items)}
        </ul>
      </nav>
      <p class="footer-note">当前主链是单机 LIO / VIO + Diff-planner + px4ctrl + mavros + multipoint。</p>
    </aside>

    <main class="main">
      <section class="hero" id="overview">
        <span class="eyebrow">Developer Portal</span>
        <h1>非凸α二开速查台</h1>
        <p class="summary">当前主链是单机 LIO / VIO + Diff-planner + px4ctrl + mavros + multipoint。D435 + VINS 是当前视觉定位主链假设，Elastic 不属于当前实机默认部署链路。这个页面先给你“先看哪里”，再把完整文档、接口、参数、部署边界和素材缺口全部接进去。</p>
        <div class="card-grid">
          {''.join(f'<article class="metric-card"><strong>{html.escape(title)}</strong><p>{html.escape(text)}</p></article>' for title, text in machine_cards)}
        </div>
        {downloads_html}
      </section>

      <section class="section" id="scenes">
        <div class="section-header">
          <div>
            <h2>场景入口</h2>
            <p>如果你已经知道自己现在要做什么，就从这里进，不要先扎进原始表格。每张卡都对应一类开发或排障任务。</p>
          </div>
        </div>
        <p class="scene-helper">点击任一场景卡片后，会自动把下方接口检索区切到对应关键词。</p>
        <div class="card-grid">{scenario_html}</div>
      </section>

      <section class="section" id="deployed">
        <div class="section-header">
          <div>
            <h2>当前已部署</h2>
            <p>这里列出的才是当前机器应进入主流程的能力。扩展能力保留在下面单独展示，不再和当前主链混写。</p>
          </div>
        </div>
        <div class="card-grid">{deployed_html}</div>
        <div class="section-header" style="margin-top:26px;">
          <div>
            <h2>扩展边界</h2>
            <p>这些能力可以进路线图，但不应写成这台机器“当前默认已落地”的能力。</p>
          </div>
        </div>
        <div class="card-grid">{extension_html}</div>
      </section>

      <section class="section" id="chains">
        <div class="section-header">
          <div>
            <h2>启动链速览</h2>
            <p>这部分不是“系统架构大图”，而是告诉你当前机器真正怎么从脚本一路起到定位、规划、控制和任务层。后续排障时先按这条链切。</p>
          </div>
        </div>
        <div class="stack-grid">{startup_html}</div>
      </section>

      <section class="section" id="params">
        <div class="section-header">
          <div>
            <h2>参数聚焦</h2>
            <p>你后续二开最常改、也最容易误改的入口先集中在这里。先动任务参数，再动规划参数，最后才考虑控制器细节。</p>
          </div>
        </div>
        <div class="stack-grid">{parameter_html}</div>
      </section>

      <section class="section" id="trace">
        <div class="section-header">
          <div>
            <h2>源码追踪</h2>
            <p>点任意接口卡或参数卡，我直接把它追到证据 ID、源码路径和建议改造层级。这样你不需要再在 CSV、知识库和源码快照之间来回跳。</p>
          </div>
        </div>
        <div class="trace-layout">
          <article class="trace-panel">
            <div class="card-kicker" id="trace-kind">{html.escape(default_trace_payload['kind'])}</div>
            <h3 id="trace-title">{html.escape(default_trace_payload['title'])}</h3>
            <p class="lead" id="trace-purpose">{html.escape(default_trace_payload['purpose'])}</p>
            <div class="trace-grid">
              <div><span>建议改造层级</span><strong id="trace-layer">{html.escape(default_trace_payload['layer'])}</strong></div>
              <div><span>证据 ID</span><strong id="trace-evidence">{html.escape(default_trace_payload['evidence'])}</strong></div>
              <div><span>机上状态</span><strong id="trace-status">{html.escape(default_trace_payload['status'])}</strong></div>
              <div><span>前置条件</span><strong id="trace-preconditions">{html.escape(default_trace_payload['preconditions'])}</strong></div>
            </div>
            <div class="trace-source" style="margin-top:12px;">
              <span>源码 / 文档路径</span>
              <code id="trace-source">{html.escape(default_trace_payload['source'])}</code>
            </div>
            <div class="trace-io-grid">
              <div class="trace-io"><span>证据结论</span><strong id="trace-claim">{html.escape(default_trace_payload['claim'])}</strong></div>
              <div class="trace-io"><span>上游 / 下游</span><strong id="trace-updown">{html.escape(default_trace_payload['upstream'])} → {html.escape(default_trace_payload['downstream'])}</strong></div>
            </div>
            <p class="card-note"><strong>示例：</strong><span id="trace-example">{html.escape(default_trace_payload['example'])}</span></p>
            <p class="card-note danger"><strong>风险：</strong><span id="trace-risk">{html.escape(default_trace_payload['risk'])}</span></p>
            <p class="footer-note" id="trace-notes">{html.escape(default_trace_payload['notes'])}</p>
          </article>
          <div class="trace-guide">
            <article class="trace-guide-card">
              <div class="card-kicker">使用顺序</div>
              <h3>先定层，再改代码</h3>
              <p>状态接口先读，动作接口先接，参数入口先改，RC 只当安全边界。你如果一上来就碰控制器，基本是在扩大问题面。</p>
            </article>
            <article class="trace-guide-card">
              <div class="card-kicker">追踪逻辑</div>
              <h3>接口、证据、源码三件事必须闭环</h3>
              <p>一个可改入口至少要同时知道它的用途、证据出处和实际源码位置。只有名字，没有路径；或者只有 topic，没有证据，都是没吃透。</p>
            </article>
            <article class="trace-guide-card">
              <div class="card-kicker">建议</div>
              <h3>第一轮二开别越层</h3>
              <p>外部系统先复用动作原语，任务改造先动 multipoint 参数，视觉或规划实验再往 `run_exp_single_*` 走，最后才考虑 `px4ctrl`。</p>
            </article>
          </div>
        </div>
      </section>

      <section class="section" id="recommended">
        <div class="section-header">
          <div>
            <h2>推荐接口</h2>
            <p>先把这几条状态订阅和动作原语接起来，建立一个不会直接碰到底层控制器的最小闭环。</p>
          </div>
        </div>
        <div class="card-grid">{recommended_html}</div>
      </section>

      <section class="section" id="explorer">
        <div class="section-header">
          <div>
            <h2>接口检索</h2>
            <p>这里保留所有接口明细，但先给你检索和分类入口，不再逼你从 CSV 原表里横向扫 16 列。</p>
          </div>
        </div>
        <div class="scene-status" id="scene-status">
          <div>
            <strong id="scene-status-title">当前未锁定场景</strong>
            <p id="scene-status-desc">你可以手动搜索，也可以点上面的场景卡，把接口区切到你当前真正关心的入口。</p>
          </div>
          <span class="scene-status-count" id="scene-status-count"></span>
        </div>
        <div class="controls">
          <input id="interface-search" type="search" placeholder="搜 topic、脚本、参数文件、用途或风险">
          <div class="filter-buttons" id="filter-buttons">
            <button class="active" data-filter="all">全部</button>
            <button data-filter="state">状态</button>
            <button data-filter="action">动作</button>
            <button data-filter="script">脚本</button>
            <button data-filter="parameter">参数</button>
            <button data-filter="rc">RC</button>
          </div>
        </div>
        <div class="interface-list" id="interface-list">{explorer_cards}</div>
      </section>

      <section class="section" id="docs">
        <div class="section-header">
          <div>
            <h2>完整文档</h2>
            <p>这里把开发主文档、截图附录和导航页直接铺开。你既可以先看上面的卡片式入口，也可以在这里顺着完整文档读。</p>
          </div>
        </div>
        <div class="doc-columns">
          <article class="doc-card">{markdown_to_html(guide_markdown)}</article>
          <article class="doc-card">{markdown_to_html(appendix_markdown)}</article>
          <article class="doc-card">{markdown_to_html(navigation_markdown)}</article>
        </div>
      </section>

      <section class="section" id="gaps">
        <div class="section-header">
          <div>
            <h2>素材缺口</h2>
            <p>这部分服务于后续补日常使用截图、补培训材料和做更强速查版。P0 先补，P1 次之。</p>
          </div>
        </div>
        <div class="table-wrap">
          <table>
            <thead>
              <tr>
                <th>状态场景</th>
                <th>需要补什么</th>
                <th>为什么缺</th>
                <th>补采指令</th>
                <th>优先级</th>
              </tr>
            </thead>
            <tbody>{gap_html}</tbody>
          </table>
        </div>
      </section>

      <section class="section" id="evidence">
        <div class="section-header">
          <div>
            <h2>证据索引</h2>
            <p>你如果要追源码、追知识库、追某条结论的出处，这里是最短入口。</p>
          </div>
        </div>
        <div class="table-wrap">
          <table>
            <thead>
              <tr>
                <th>证据 ID</th>
                <th>结论</th>
                <th>来源路径</th>
              </tr>
            </thead>
            <tbody>{evidence_html}</tbody>
          </table>
        </div>
      </section>

      <section class="section" id="downloads">
        <div class="section-header">
          <div>
            <h2>导出件</h2>
            <p>HTML 负责读起来舒服；Word 和 Excel 继续保留给归档、转发和二次加工。</p>
          </div>
        </div>
        {downloads_html}
        <p class="footer-note">如果你更习惯线性阅读，直接看 <code>开发文档包.docx</code>；如果要快速接机器，看 <code>接口总表.xlsx</code> 的 <code>场景速查</code> 和 <code>推荐接入</code>。</p>
      </section>
    </main>
  </div>

  <script>
    const searchInput = document.getElementById('interface-search');
    const buttons = Array.from(document.querySelectorAll('#filter-buttons button'));
    const cards = Array.from(document.querySelectorAll('.interface-row'));
    const sceneCards = Array.from(document.querySelectorAll('.scene-jump'));
    const traceTriggers = Array.from(document.querySelectorAll('.trace-trigger'));
    const sceneStatusTitle = document.getElementById('scene-status-title');
    const sceneStatusDesc = document.getElementById('scene-status-desc');
    const sceneStatusCount = document.getElementById('scene-status-count');
    const traceKind = document.getElementById('trace-kind');
    const traceTitle = document.getElementById('trace-title');
    const tracePurpose = document.getElementById('trace-purpose');
    const traceLayer = document.getElementById('trace-layer');
    const traceEvidence = document.getElementById('trace-evidence');
    const traceStatus = document.getElementById('trace-status');
    const tracePreconditions = document.getElementById('trace-preconditions');
    const traceSource = document.getElementById('trace-source');
    const traceClaim = document.getElementById('trace-claim');
    const traceUpdown = document.getElementById('trace-updown');
    const traceExample = document.getElementById('trace-example');
    const traceRisk = document.getElementById('trace-risk');
    const traceNotes = document.getElementById('trace-notes');
    let activeFilter = 'all';

    function setActiveFilter(filter) {{
      activeFilter = filter;
      buttons.forEach((item) => item.classList.toggle('active', item.dataset.filter === filter));
    }}

    function applyFilters() {{
      const tokens = searchInput.value.trim().toLowerCase().split(/\\s+/).filter(Boolean);
      let visibleCount = 0;
      cards.forEach((card) => {{
        const category = card.dataset.category;
        const haystack = card.dataset.search.toLowerCase();
        const matchesFilter = activeFilter === 'all' || category === activeFilter;
        const matchesKeyword = !tokens.length || tokens.every((token) => haystack.includes(token));
        const visible = matchesFilter && matchesKeyword;
        card.style.display = visible ? '' : 'none';
        if (visible) {{
          visibleCount += 1;
        }}
      }});
      sceneStatusCount.textContent = '当前显示 ' + visibleCount + ' 条接口';
    }}

    buttons.forEach((button) => {{
      button.addEventListener('click', () => {{
        setActiveFilter(button.dataset.filter);
        sceneCards.forEach((card) => card.classList.remove('active-scene'));
        sceneStatusTitle.textContent = '当前未锁定场景';
        sceneStatusDesc.textContent = '你可以手动搜索，也可以点上面的场景卡，把接口区切到你当前真正关心的入口。';
        applyFilters();
      }});
    }});

    sceneCards.forEach((card) => {{
      function jumpToScene() {{
        const query = card.dataset.sceneQuery || '';
        const filter = card.dataset.sceneFilter || 'all';
        const title = card.dataset.sceneTitle || '场景入口';
        const summary = card.dataset.sceneSummary || '';

        sceneCards.forEach((item) => item.classList.toggle('active-scene', item === card));
        searchInput.value = query;
        setActiveFilter(filter);
        sceneStatusTitle.textContent = '当前场景：' + title;
        sceneStatusDesc.textContent = summary;
        applyFilters();
        document.getElementById('explorer').scrollIntoView({{ behavior: 'smooth', block: 'start' }});
      }}

      card.addEventListener('click', jumpToScene);
      card.addEventListener('keydown', (event) => {{
        if (event.key === 'Enter' || event.key === ' ') {{
          event.preventDefault();
          jumpToScene();
        }}
      }});
    }});

    function updateTrace(trigger, shouldScroll) {{
      traceTriggers.forEach((item) => item.classList.toggle('active-trace', item === trigger));
      traceKind.textContent = trigger.dataset.traceKind || '';
      traceTitle.textContent = trigger.dataset.traceTitle || '';
      tracePurpose.textContent = trigger.dataset.tracePurpose || '';
      traceLayer.textContent = trigger.dataset.traceLayer || '';
      traceEvidence.textContent = trigger.dataset.traceEvidence || '';
      traceStatus.textContent = trigger.dataset.traceStatus || '';
      tracePreconditions.textContent = trigger.dataset.tracePreconditions || '';
      traceSource.textContent = trigger.dataset.traceSource || '';
      traceClaim.textContent = trigger.dataset.traceClaim || '';
      traceUpdown.textContent = (trigger.dataset.traceUpstream || '') + ' → ' + (trigger.dataset.traceDownstream || '');
      traceExample.textContent = trigger.dataset.traceExample || '';
      traceRisk.textContent = trigger.dataset.traceRisk || '';
      traceNotes.textContent = trigger.dataset.traceNotes || '';

      if (shouldScroll) {{
        document.getElementById('trace').scrollIntoView({{ behavior: 'smooth', block: 'start' }});
      }}
    }}

    traceTriggers.forEach((trigger) => {{
      function activateTrace() {{
        const shouldScroll = !trigger.closest('#explorer');
        updateTrace(trigger, shouldScroll);
      }}

      trigger.addEventListener('click', activateTrace);
      trigger.addEventListener('keydown', (event) => {{
        if (event.key === 'Enter' || event.key === ' ') {{
          event.preventDefault();
          activateTrace();
        }}
      }});
    }});

    searchInput.addEventListener('input', applyFilters);
    applyFilters();
    const defaultTrace = traceTriggers.find((item) => item.dataset.traceDefault === 'true') || traceTriggers[0];
    if (defaultTrace) {{
      updateTrace(defaultTrace, false);
    }}
  </script>
</body>
</html>
"""

    out_path.write_text(html_text, encoding="utf-8")


def markdown_bundle_to_docx(markdown_paths: list[Path], out_path: Path) -> None:
    ensure_docx_import()
    from docx import Document

    document = Document()
    for markdown_path in markdown_paths:
        for raw_line in markdown_path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line:
                document.add_paragraph("")
                continue
            if line.startswith("# "):
                document.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                document.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                document.add_heading(line[4:], level=3)
            elif line.startswith("- "):
                document.add_paragraph(line[2:], style="List Bullet")
            else:
                document.add_paragraph(line)
    document.save(out_path)


def export_xlsx_bundle() -> None:
    node_modules = node_modules_path()
    link_path, created = ensure_node_modules_link(node_modules)
    payload_path = SCRIPT_ROOT / ".uav_dev_docs_export_data.json"
    env = os.environ.copy()
    if node_modules is not None:
        env["NODE_PATH"] = str(node_modules)
    try:
        payload_path.write_text(
            json.dumps(build_export_payload(), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        subprocess.run(
            [
                str(node_executable()),
                str(SCRIPT_ROOT / "export_uav_dev_docs_xlsx.mjs"),
                str(payload_path),
            ],
            cwd=ROOT,
            env=env,
            check=True,
        )
    finally:
        payload_path.unlink(missing_ok=True)
        if created and link_path is not None and link_path.exists():
            link_path.unlink()


def main() -> None:
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    export_xlsx_bundle()
    docx_path = OUT_ROOT / "非凸α-开发者文档包.docx"
    html_path = OUT_ROOT / "非凸α-二开速查台.html"
    markdown_bundle_to_docx(
        [
            DOC_ROOT / "01-developer-guide.md",
            DOC_ROOT / "04-screenshot-state-appendix.md",
            DOC_ROOT / "05-developer-navigation.md",
        ],
        docx_path,
    )
    build_html_portal(build_export_payload(), html_path)
    print("WROTE", docx_path)
    print("WROTE", html_path)


if __name__ == "__main__":
    main()
